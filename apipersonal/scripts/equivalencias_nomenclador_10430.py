"""
Genera el Excel de equivalencias ocupacion -> cargo del nomenclador 10430 para que lo
valide el usuario antes de asignar cargos (tablas de la mig 094).

Por cada agente 10430 activo toma agrupamiento y categoria del ultimo registro del
SIAPE (tabla historial) y la ocupacion de agentes. La propuesta es por
(ocupacion, agrupamiento SIAPE) -> familia de cargos (codigos del mismo agrupamiento).

Reglas de la asignacion:
  - grado y clase salen de la categoria (nomenclador_categorias)
  - clase 4 -> cargo Inicial "E" (codigo 0000) en cualquier agrupamiento
  - clase A/1/2/3 -> letra A/B/C/D; se busca en la familia el codigo con esa letra
  - si la familia no tiene esa letra -> <agrup>-0000-<grado>-<clase> (cargo no
    contemplado, nota del Buscador del nomenclador, ej. Chofer "A")

Uso:  python equivalencias_nomenclador_10430.py "<salida.xlsx>"
"""
import os
import sys
from collections import defaultdict

import pymysql
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# DB_* del .env del API (sin pisar lo que ya venga en el entorno)
_env = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".env")
if os.path.exists(_env):
    for _l in open(_env, encoding="utf-8", errors="ignore"):
        if _l.startswith("DB_") and "=" in _l:
            _k, _v = _l.strip().split("=", 1)
            os.environ.setdefault(_k, _v)

AGRUP_ID = {"SERVICIO": 1, "OBRERO": 2, "ADMINISTRATIVO": 3, "TECNICO": 4, "PROFESIONAL": 5}
LETRA_DE_CLASE = {"A": "A", "1": "B", "2": "C", "3": "D"}

# Familias: agrupamiento -> codigos (una por letra)
F = {
    "ENFERMERO": (4, [241, 242, 243, 244]),
    "AUX_ELECTRO": (4, [65, 66, 67, 68]),
    "AUX_FARMACIA": (4, [77, 78, 79, 80]),
    "AUX_HEMOTERAPIA": (4, [89, 90, 91, 92]),
    "AYU_HEMOTERAPIA": (4, [169, 170, 171, 172]),
    "AUX_RADIOLOGIA": (4, [133, 134, 135, 136]),
    "ESTERILIZACION": (4, [57, 58, 59, 60]),
    "INSTRUMENTADOR": (4, [309, 310, 311, 312]),
    "TEC_RADIOLOGIA": (4, [529, 530, 531, 532]),
    "TEC_MANTENIMIENTO": (4, [505, 506, 507, 508]),
    "HISTOLOGIA": (4, [413, 414, 415, 416]),
    "LABORATORISTA": (4, [573, 574, 575, 576]),
    "CITOLOGIA": (4, [193, 194, 195, 196]),
    "ANALISTA_ADM_TEC": (4, [5, 6, 7, 8]),
    "ADMINISTRATIVO": (3, [1, 2, 3, 4]),
    "CAMILLERO": (1, [3]),
    "GUARDIAN": (1, [45, 46, 22, 23]),
    "MUCAMA": (1, [36]),
    "CHOFER": (1, [13, 14, 15]),
    "COCINA": (1, [8, 9, 10]),
    "ORDENANZA": (1, [37, 38]),
    "PEON": (1, [40]),
    "MANT_AYUDANTE": (2, [119]),
    "ELECTRICISTA": (2, [58, 59]),
    "CONMUTADOR": (2, [40, 41]),
    "AYU_OFICIOS": (2, [18]),
    "ABOGADO": (5, [1, 2, 3, 4]),
    "PROF_ENFERMERIA": (5, [382, 383, 384, 385]),
    "PROF_BIOIMAGENES": (5, [386, 387, 388, 389]),
}

# (ocupacion normalizada, agrupamiento SIAPE) -> (familia, confianza, motivo)
ENF = ("ENFERMERO", "ALTA", "")
PROP = {
    ("ENFERMERA PROFESIONAL", "TECNICO"): ENF,
    ("ENFERMERO PROFESIONAL", "TECNICO"): ENF,
    ("ENFERMERA", "TECNICO"): ENF,
    ("ENFERMERO", "TECNICO"): ENF,
    ("ENFERMERA A", "TECNICO"): ENF,
    ("ENFERMERA B", "TECNICO"): ENF,
    ("ENFERMERA C", "TECNICO"): ENF,
    ("ENFERMERO B", "TECNICO"): ENF,
    ("ENFERMERO C", "TECNICO"): ENF,
    ("ENF. UNIVERSITARIA", "TECNICO"): ("ENFERMERO", "MEDIA", "Universitaria pero el SIAPE la tiene en Técnico"),
    ("ENFERMERA UNIVERSITARIA", "TECNICO"): ("ENFERMERO", "MEDIA", "Universitaria pero el SIAPE la tiene en Técnico"),
    ("LICENCIADO EN ENFERMERIA", "TECNICO"): ("ENFERMERO", "MEDIA", "Licenciado pero el SIAPE lo tiene en Técnico"),
    ("AUXILIAR DE ELECTROENCEFALOCARDIOGRAFIA", "TECNICO"): ("AUX_ELECTRO", "ALTA", ""),
    ("AUXILIAR DE FARMACIA", "TECNICO"): ("AUX_FARMACIA", "ALTA", ""),
    ("AUXILIAR FARMACIA", "TECNICO"): ("AUX_FARMACIA", "ALTA", ""),
    ("AUXILIAR DE FARMACIA A", "TECNICO"): ("AUX_FARMACIA", "ALTA", ""),
    ("AUXILIAR DE FARMACIA B", "TECNICO"): ("AUX_FARMACIA", "ALTA", ""),
    ("AUXILIAR DE HEMOTERAPIA", "TECNICO"): ("AUX_HEMOTERAPIA", "ALTA", ""),
    ("AYUDANTE DE HEMOTERAPIA", "TECNICO"): ("AYU_HEMOTERAPIA", "ALTA", ""),
    ("AUXILIAR DE RADIOLOGÍA", "TECNICO"): ("AUX_RADIOLOGIA", "ALTA", ""),
    ("AYUDANTE DE ESTERILIZACION INICIAL", "TECNICO"): ("ESTERILIZACION", "ALTA", "La familia cambia de nombre por letra (Supervisor A … Ayudante D)"),
    ("INSTRUMENTADOR QUIRURGICO", "TECNICO"): ("INSTRUMENTADOR", "ALTA", ""),
    ("TECNICO EN INSTRUMENTACION QUIRURGICA", "TECNICO"): ("INSTRUMENTADOR", "ALTA", ""),
    ("TECNICO EN RADIOLOGIA Y RADIOTERAPIA", "TECNICO"): ("TEC_RADIOLOGIA", "ALTA", ""),
    ("LICENCIADO EN PRODUCCION EN BIOIMAGENES", "TECNICO"): ("TEC_RADIOLOGIA", "MEDIA", "Licenciado pero el SIAPE lo tiene en Técnico"),
    ("PREPARADOR DE HISTOLOGÍA B", "TECNICO"): ("HISTOLOGIA", "ALTA", ""),
    ("TÉCNICO LABORATORISTA C", "TECNICO"): ("LABORATORISTA", "ALTA", ""),
    ("TÉCNICO LABORATORISTA D", "TECNICO"): ("LABORATORISTA", "ALTA", ""),
    ("TEC. QUIMICO", "TECNICO"): ("LABORATORISTA", "MEDIA", "No hay 'Técnico Químico' en el nomenclador"),
    ("TEC.CITOLOGA", "TECNICO"): ("CITOLOGIA", "ALTA", ""),
    ("OPERARIO", "TECNICO"): ("TEC_MANTENIMIENTO", "BAJA", "Operario revistando en Técnico: confirmar tarea"),
    ("AYUDANTE ADMINISTRATIVA", "TECNICO"): ("ANALISTA_ADM_TEC", "BAJA", "Administrativa revistando en Técnico: confirmar"),
    ("AUXILIAR ADMINISTRATIVO", "ADMINISTRATIVO"): ("ADMINISTRATIVO", "ALTA", ""),
    ("OFICIAL ADMINISTRATIVO", "ADMINISTRATIVO"): ("ADMINISTRATIVO", "ALTA", ""),
    ("AYUDANTE ADMINISTRATIVA", "ADMINISTRATIVO"): ("ADMINISTRATIVO", "ALTA", ""),
    ("AYUDANTE ADMINISTRATIVO", "ADMINISTRATIVO"): ("ADMINISTRATIVO", "ALTA", ""),
    ("AYUDANTE ADMINISTRATIVO INICIAL", "ADMINISTRATIVO"): ("ADMINISTRATIVO", "ALTA", ""),
    ("AUXILIAR DE FARMACIA", "ADMINISTRATIVO"): ("ADMINISTRATIVO", "BAJA", "Trabaja en farmacia pero revista en Administrativo"),
    ("AUXILIAR DE FARMACIA B", "ADMINISTRATIVO"): ("ADMINISTRATIVO", "BAJA", "Trabaja en farmacia pero revista en Administrativo"),
    ("MUCAMA", "ADMINISTRATIVO"): ("ADMINISTRATIVO", "BAJA", "Mucama revistando en Administrativo"),
    ("CAMILLERO", "SERVICIO"): ("CAMILLERO", "ALTA", "Solo existe Camillero \"D\""),
    ("GUARDIAN", "SERVICIO"): ("GUARDIAN", "ALTA", "La clase C es \"Guardián Encargado C\""),
    ("MUCAMA", "SERVICIO"): ("MUCAMA", "ALTA", "Solo existe Mucama \"D\""),
    ("MUCAMO", "SERVICIO"): ("MUCAMA", "ALTA", "Solo existe Mucama \"D\""),
    ("CHOFER A", "SERVICIO"): ("CHOFER", "ALTA", "No existe Chofer \"A\": va 0000 (ejemplo del propio nomenclador)"),
    ("CHOFER B", "SERVICIO"): ("CHOFER", "ALTA", ""),
    ("CARNICERO", "SERVICIO"): ("COCINA", "BAJA", "Carnicero solo existe en Obrero; en Servicio lo más parecido es Cocina"),
    ("AYUDANTE ADMINISTRATIVA", "SERVICIO"): ("ORDENANZA", "BAJA", "Administrativa revistando en Servicio: confirmar"),
    ("AYUDANTE ADMINISTRATIVO", "SERVICIO"): ("ORDENANZA", "BAJA", "Administrativo revistando en Servicio: confirmar"),
    ("AYUDANTE DE MANTENIMIENTO", "SERVICIO"): ("PEON", "BAJA", "Mantenimiento revistando en Servicio"),
    ("OPERARIO", "SERVICIO"): ("PEON", "MEDIA", ""),
    ("AYUDANTE DE MANTENIMIENTO", "OBRERO"): ("MANT_AYUDANTE", "ALTA", "Solo existe \"D\""),
    ("ELECTRICISTA INSTALADOR MEDIO OFICIAL", "OBRERO"): ("ELECTRICISTA", "ALTA", ""),
    ("OPERADOR DE CONMUTADOR TELEFONICO", "OBRERO"): ("CONMUTADOR", "ALTA", ""),
    ("OPERARIO", "OBRERO"): ("AYU_OFICIOS", "MEDIA", "No hay 'Operario' genérico; Ayudante de Oficios es solo \"D\""),
    ("ABOGADO", "PROFESIONAL"): ("ABOGADO", "ALTA", ""),
    ("ENFERMERO B", "PROFESIONAL"): ("PROF_ENFERMERIA", "BAJA", "Categoría 3 no existe en Profesional (arranca en 8)"),
    ("ENFERMERO PROFESIONAL", "PROFESIONAL"): ("PROF_ENFERMERIA", "BAJA", "Categoría 3 no existe en Profesional (arranca en 8)"),
    ("LICENCIADO EN ENFERMERIA", "PROFESIONAL"): ("PROF_ENFERMERIA", "BAJA", "Categoría 3 no existe en Profesional (arranca en 8)"),
    ("LICENCIADO EN PRODUCCION EN BIOIMAGENES", "PROFESIONAL"): ("PROF_BIOIMAGENES", "BAJA", "Categoría 3 no existe en Profesional (arranca en 8)"),
}

SQL_AGENTES = """
WITH act AS (
  SELECT a.id, a.dni, a.ocupacion_id FROM agentes a
  WHERE a.ley_id IN (1, 3) AND a.estado_empleo = 'ACTIVO' AND a.deleted_at IS NULL
), ult AS (
  SELECT h.dni, h.agrupamiento, h.categoria_salarial, h.fecha_desde,
         ROW_NUMBER() OVER (PARTITION BY h.dni ORDER BY h.fecha_desde DESC) rn
  FROM historial h WHERE h.dni IN (SELECT dni FROM act)
)
SELECT act.id, act.dni, CONCAT(p.apellido, ', ', p.nombre) nombre, o.id, o.nombre,
       u.agrupamiento, u.categoria_salarial
FROM act
JOIN personal p ON p.dni = act.dni
JOIN ocupaciones o ON o.id = act.ocupacion_id
LEFT JOIN ult u ON u.dni = act.dni AND u.rn = 1
ORDER BY p.apellido, p.nombre
"""


def norm(s):
    return " ".join((s or "").upper().split())


def conectar():
    return pymysql.connect(
        host=os.environ.get("DB_HOST", "127.0.0.1"),
        port=int(os.environ.get("DB_PORT", "3306")),
        user=os.environ.get("DB_USER", "root"),
        password=os.environ.get("DB_PASSWORD", ""),
        database=os.environ.get("DB_NAME", "personalv5"),
        charset="utf8mb4",
    )


def proponer(agentes, cargos, categorias):
    """Devuelve filas por agente con el codigo propuesto y su tipo."""
    filas = []
    for agente_id, dni, nombre, ocup_id, ocup, agrup, cat in agentes:
        agrup = norm(agrup) or None
        cat = int(cat) if cat and str(cat).isdigit() else None
        prop = PROP.get((norm(ocup), agrup))
        base = dict(agente_id=agente_id, dni=dni, nombre=nombre, ocup_id=ocup_id, ocup=ocup.strip(),
                    agrup=agrup, cat=cat, grado=None, clase=None, codigo=None, denominacion=None,
                    familia=prop[0] if prop else None, tipo=None, motivo="")
        aid = AGRUP_ID.get(agrup)
        if aid is None:
            base.update(tipo="REVISAR", motivo="Sin agrupamiento en el SIAPE" if not agrup
                        else f"Agrupamiento {agrup} no está en el nomenclador")
        elif (aid, cat) not in categorias:
            base.update(tipo="REVISAR", motivo=f"Categoría {cat} no existe en {agrup}")
        elif not prop:
            base.update(tipo="REVISAR", motivo="Sin equivalencia propuesta")
        else:
            grado, clase = categorias[(aid, cat)]
            base.update(grado=grado, clase=clase)
            fam_agr, codigos = F[prop[0]]
            if clase == "4":
                cod, den, tipo = 0, cargos[(aid, 0)][0], "INICIAL"
            else:
                letra = LETRA_DE_CLASE[clase]
                hit = [c for c in codigos if cargos[(fam_agr, c)][1] == letra]
                if hit:
                    cod, den, tipo = hit[0], cargos[(fam_agr, hit[0])][0], "EXACTO"
                else:
                    base_den = cargos[(fam_agr, codigos[0])][0].split('"')[0].strip()
                    cod, den, tipo = 0, f'{base_den} "{letra}" (no contemplado)', "NO CONTEMPLADO"
            base.update(codigo=f"{aid}-{cod:04d}-{grado}-{clase}", denominacion=den, tipo=tipo)
        filas.append(base)
    return filas


def hoja(ws, encabezados, filas, anchos):
    ws.append(encabezados)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for f in filas:
        ws.append(f)
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    if len(sys.argv) < 2:
        sys.exit('Uso: python equivalencias_nomenclador_10430.py "<salida.xlsx>"')
    cn = conectar()
    with cn.cursor() as cur:
        cur.execute(SQL_AGENTES)
        agentes = cur.fetchall()
        cur.execute("SELECT agrupamiento_id, codigo, denominacion, letra, vigente FROM nomenclador_cargos")
        todos = cur.fetchall()
        cur.execute("SELECT agrupamiento_id, categoria, grado, clase FROM nomenclador_categorias")
        categorias = {(a, c): (g, k) for a, c, g, k in cur.fetchall()}
    cn.close()
    cargos = {(a, c): (d, l) for a, c, d, l, _ in todos}
    filas = proponer(agentes, cargos, categorias)

    wb = Workbook()
    ws = wb.active
    ws.title = "Leeme"
    for linea in [
        "Equivalencias ocupación → cargo del nomenclador Ley 10430 (agentes activos)",
        "",
        "Qué hay que validar: la hoja Equivalencias. Una fila por ocupación + agrupamiento del SIAPE.",
        "  • VALIDADO: SI si la familia propuesta está bien, NO si no.",
        "  • CORRECCIÓN: si es NO, poné el agrupamiento y código(s) correctos (ver hoja Catálogo), ej. 4-241/244.",
        "",
        "Cómo se calcula el código de cada agente (hoja Agentes):",
        "  • Agrupamiento y categoría salen del último registro del SIAPE.",
        "  • Grado y clase salen de la categoría según la tabla del nomenclador.",
        "  • Clase 4 = cargo Inicial \"E\" (código 0000), en cualquier agrupamiento.",
        "  • Clase A/1/2/3 = letra A/B/C/D: se toma el cargo de la familia con esa letra.",
        "  • Si la familia no tiene esa letra (ej. Mucama solo existe \"D\"), va 0000 = cargo no contemplado,",
        "    como indica la nota del Buscador del nomenclador (ej. Chofer \"A\").",
        "  • REVISAR = no se puede calcular (sin SIAPE, agrupamiento o categoría fuera del nomenclador).",
    ]:
        ws.append([linea])
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 110

    grupos = defaultdict(list)
    for f in filas:
        grupos[(f["ocup_id"], f["ocup"], f["agrup"])].append(f)
    eq = []
    for (ocup_id, ocup, agrup), fs in sorted(grupos.items(), key=lambda kv: (-len(kv[1]), kv[0][1])):
        prop = PROP.get((norm(ocup), agrup))
        if prop:
            fam_agr, codigos = F[prop[0]]
            familia = " / ".join(cargos[(fam_agr, c)][0] for c in codigos)
            cods = f"{fam_agr}-" + ",".join(str(c) for c in codigos)
            conf, motivo = prop[1], prop[2]
        else:
            familia, cods, conf, motivo = "", "", "SIN PROPUESTA", fs[0]["motivo"]
        cats = sorted({f["cat"] for f in fs if f["cat"] is not None})
        tipos = defaultdict(int)
        for f in fs:
            tipos[f["tipo"]] += 1
        eq.append([ocup_id, ocup, agrup or "(sin SIAPE)", len(fs), ", ".join(map(str, cats)),
                   familia, cods, conf, motivo,
                   tipos["EXACTO"], tipos["INICIAL"], tipos["NO CONTEMPLADO"], tipos["REVISAR"], "", ""])
    ws = wb.create_sheet("Equivalencias")
    hoja(ws, ["Ocupación id", "Ocupación", "Agrupamiento SIAPE", "Agentes", "Categorías",
              "Cargo propuesto (familia)", "Códigos", "Confianza", "Motivo / observación",
              "Exacto", "Inicial", "No contemplado", "Revisar", "VALIDADO", "CORRECCIÓN"],
         eq, [10, 38, 16, 9, 18, 60, 16, 14, 50, 8, 8, 10, 8, 11, 30])
    dv = DataValidation(type="list", formula1='"SI,NO"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"N2:N{len(eq) + 1}")
    colores = {"ALTA": "C6EFCE", "MEDIA": "FFEB9C", "BAJA": "F8CBAD", "SIN PROPUESTA": "FF9999"}
    for row in ws.iter_rows(min_row=2):
        row[7].fill = PatternFill("solid", fgColor=colores.get(row[7].value, "FFFFFF"))
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("Agentes")
    hoja(ws, ["DNI", "Apellido y nombre", "Ocupación", "Agrupamiento SIAPE", "Categoría", "Grado", "Clase",
              "Código propuesto", "Denominación propuesta", "Tipo", "Motivo"],
         [[f["dni"], f["nombre"], f["ocup"], f["agrup"], f["cat"], f["grado"], f["clase"],
           f["codigo"], f["denominacion"], f["tipo"], f["motivo"]] for f in filas],
         [11, 32, 34, 16, 10, 7, 6, 18, 48, 16, 40])

    ws = wb.create_sheet("Catálogo")
    nombres = {v: k for k, v in AGRUP_ID.items()}
    hoja(ws, ["Agrupamiento", "Código", "Denominación", "Letra", "Vigente"],
         [[f"{a} {nombres[a]}", c, d, l, "SI" if v else "NO"] for a, c, d, l, v in sorted(todos)],
         [18, 9, 70, 7, 8])

    wb.save(sys.argv[1])
    tot = defaultdict(int)
    for f in filas:
        tot[f["tipo"]] += 1
    print(f"{len(filas)} agentes, {len(eq)} equivalencias -> {dict(tot)}")
    print("sin propuesta:", [(r[1], r[2], r[3]) for r in eq if r[7] == "SIN PROPUESTA"])


if __name__ == "__main__":
    main()
