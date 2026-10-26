"""
Carga el nomenclador de cargos de la Ley 10430 en personalv5 (tablas de la mig 094):
  - nomenclador_categorias: categoria -> grado/clase por agrupamiento (hoja "data")
  - nomenclador_cargos: cargos de las hojas Servicio/Obrero/Administrativo/Tecnico/Profesional

Fuente: "AUTO NOMENCLADOR (AlfaV1.09).xlsx" (EX-2024-40323859-GDEBA-DSTECMTRAGP).

Uso:  python importar_nomenclador_10430.py "D:\\G\\AUTO NOMENCLADOR (AlfaV1.09).xlsx" [--dry-run]
Re-ejecutable: actualiza por (agrupamiento, codigo) y (agrupamiento, categoria), no borra
filas (agentes_cargo_nomenclador las referencia por FK).
"""
import os
import re
import sys

import openpyxl
import pymysql

# DB_* del .env del API (sin pisar lo que ya venga en el entorno)
_env = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".env")
if os.path.exists(_env):
    for _l in open(_env, encoding="utf-8", errors="ignore"):
        if _l.startswith("DB_") and "=" in _l:
            _k, _v = _l.strip().split("=", 1)
            os.environ.setdefault(_k, _v)

# hoja -> agrupamiento_id (numero oficial del nomenclador)
HOJAS = {"Servicio": 1, "Obrero": 2, "Administrativo": 3, "Técnico": 4, "Profesional": 5}

# hoja "data": columna (0-based) donde arranca el bloque categoria|grado|clase de cada agrupamiento
BLOQUES_CATEGORIA = {1: 15, 2: 18, 3: 21, 4: 24, 5: 27}

ROMANO = re.compile(r"^(XI{0,3}|XIV|IX|IV|VI{0,3}|I{1,3})$")
CLASES = {"A", "1", "2", "3", "4"}
# letra entre comillas al final de la denominacion: "B", “B”, «B»
LETRA_EN_NOMBRE = re.compile(r"[\"“”«»'´`]\s*([A-E])\s*[\"“”«»'´`]")


def txt(v, n=None):
    if v is None:
        return None
    s = str(v).strip()
    if not s or s.startswith("#"):  # #VALUE!, #N/A de la planilla
        return None
    return s[:n] if n else s


def num(v):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return int(f) if f == int(f) else None


def leer_categorias(wb):
    filas = []
    for r in wb["data"].iter_rows(min_row=3, values_only=True):
        for agr, col in BLOQUES_CATEGORIA.items():
            if len(r) <= col + 2:
                continue
            cat, grado, clase = num(r[col]), txt(r[col + 1]), txt(r[col + 2])
            if cat is None or not grado or not ROMANO.match(grado):
                continue
            clase = clase.split(".")[0] if clase else None  # "3.0" -> "3"
            if clase not in CLASES:
                continue
            filas.append((agr, cat, grado, clase))
    return filas


def leer_cargos(wb):
    filas = []
    for hoja, agr in HOJAS.items():
        rows = list(wb[hoja].iter_rows(values_only=True))
        inicio = next(i for i, r in enumerate(rows) if r and r[0] == "CODIGO") + 1
        for r in rows[inicio:]:
            if not r or num(r[0]) is None:
                continue
            r = list(r) + [None] * (10 - len(r))
            denominacion = txt(r[1], 255)
            letra = txt(r[2])
            letra = letra.upper()[:1] if letra and letra.upper()[:1] in "ABCDE" else None
            if letra is None and denominacion:
                m = LETRA_EN_NOMBRE.search(denominacion)
                letra = m.group(1) if m else None
            dejado = txt(r[7], 255)
            filas.append((
                agr, num(r[0]), denominacion, letra,
                txt(r[3]), txt(r[4]), txt(r[5]),
                txt(r[6], 255), dejado, txt(r[9], 255),
                0 if dejado else 1,
            ))
    return filas


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args:
        sys.exit('Uso: python importar_nomenclador_10430.py "<ruta al xlsx>" [--dry-run]')
    dry = "--dry-run" in sys.argv
    wb = openpyxl.load_workbook(args[0], data_only=True, read_only=True)

    categorias = leer_categorias(wb)
    cargos = leer_cargos(wb)
    for agr in sorted(HOJAS.values()):
        cats = sorted(c[1] for c in categorias if c[0] == agr)
        n = sum(1 for c in cargos if c[0] == agr)
        print(f"agrupamiento {agr}: categorias {cats[0]}-{cats[-1]} ({len(cats)}), cargos {n}")
    sin_letra = [(c[0], c[1], c[2]) for c in cargos if c[3] is None]
    print(f"cargos sin letra: {len(sin_letra)}", sin_letra[:10])
    print(f"cargos dejados sin efecto: {sum(1 for c in cargos if not c[10])}")
    if dry:
        print("--dry-run: no se escribio nada")
        return

    cn = pymysql.connect(
        host=os.environ.get("DB_HOST", "127.0.0.1"),
        port=int(os.environ.get("DB_PORT", "3306")),
        user=os.environ.get("DB_USER", "root"),
        password=os.environ.get("DB_PASSWORD", ""),
        database=os.environ.get("DB_NAME", "personalv5"),
        charset="utf8mb4",
    )
    try:
        with cn.cursor() as cur:
            cur.executemany(
                """INSERT INTO nomenclador_categorias (agrupamiento_id, categoria, grado, clase)
                   VALUES (%s, %s, %s, %s)
                   ON DUPLICATE KEY UPDATE grado = VALUES(grado), clase = VALUES(clase)""",
                categorias,
            )
            cur.executemany(
                """INSERT INTO nomenclador_cargos
                     (agrupamiento_id, codigo, denominacion, letra, descripcion_tareas, requisitos,
                      caracteristicas, incorporado_por, dejado_sin_efecto_por, ambito_aplicacion, vigente)
                   VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                   ON DUPLICATE KEY UPDATE
                     denominacion = VALUES(denominacion), letra = VALUES(letra),
                     descripcion_tareas = VALUES(descripcion_tareas), requisitos = VALUES(requisitos),
                     caracteristicas = VALUES(caracteristicas), incorporado_por = VALUES(incorporado_por),
                     dejado_sin_efecto_por = VALUES(dejado_sin_efecto_por),
                     ambito_aplicacion = VALUES(ambito_aplicacion), vigente = VALUES(vigente)""",
                cargos,
            )
        cn.commit()
        print(f"OK: {len(categorias)} categorias, {len(cargos)} cargos")
    finally:
        cn.close()


if __name__ == "__main__":
    main()
